"""Excel-шаблон и импорт плана уроков. Один формат — тот же LessonPlanItem."""

from __future__ import annotations

import io
import logging
import re
from datetime import date, datetime
from typing import Any

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from rest_framework import status

from .choices import PlanItemStatus
from .files_storage import content_disposition, sanitize_filename
from .models import LessonPlan, LessonPlanItem, ScheduleEvent
from .plan_dates import (
    INTERVAL_LABELS,
    apply_sequence_dates,
    generate_plan_dates,
    iso_plan_date,
    normalize_interval,
    parse_plan_date,
)
from .plan_levels import get_plan_level_label, get_plan_level_options, normalize_plan_level_id
from .plan_subjects import get_plan_subject_label, get_plan_subject_options, normalize_plan_subject_id

logger = logging.getLogger("cabinet.plan_excel")

TEMPLATE_VERSION = "lesson_plan_v2"
SUPPORTED_TEMPLATE_VERSIONS = frozenset({"", "lesson_plan_v1", "lesson_plan_v2"})
MAX_FILE_BYTES = 4 * 1024 * 1024
MAX_LESSONS = 500
MAX_TITLE = 255
MAX_TOPIC = 500
MAX_SUBTOPIC = 255
MAX_TASK_NUMBER = 32
TEMPLATE_SPARE_ROWS = 80

SHEET_LESSONS = "Уроки"
SHEET_INSTRUCTIONS = "Инструкция"
SHEET_SETTINGS = "Настройки"
SHEET_META = "_meta"

MODE_REPLACE_ALL = "replace_all"
MODE_INSERT = "insert"
MODE_REPLACE_DATES = "replace_dates"
IMPORT_MODES = (MODE_REPLACE_ALL, MODE_INSERT, MODE_REPLACE_DATES)
LEGACY_MODE_MAP = {
    "replace": MODE_REPLACE_ALL,
    "append": MODE_INSERT,
}

GRADE_CHOICES = ("5", "6", "7", "8", "9", "10", "11", "10–11")

HEADER_FILL = PatternFill("solid", fgColor="334155")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
CELL_FONT = Font(name="Calibri", size=11)
HINT_FONT = Font(name="Calibri", size=10, color="64748B")
TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="0F172A")
SECTION_FONT = Font(name="Calibri", bold=True, size=12, color="0F172A")
LABEL_FONT = Font(name="Calibri", bold=True, size=11, color="0F172A")
BOX_FILL = PatternFill("solid", fgColor="F8FAFC")
SECTION_FILL = PatternFill("solid", fgColor="EEF2F7")
WARN_FILL = PatternFill("solid", fgColor="FFFBEB")
BAND_FILL = PatternFill("solid", fgColor="F8FAFC")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
THIN = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(vertical="center", wrap_text=True)
TOP = Alignment(vertical="center")

COLUMNS = (
    {"key": "number", "title": "№", "width": 6, "locked": True, "hidden": False, "wrap": False},
    {"key": "item_id", "title": "ID урока", "width": 12, "locked": True, "hidden": True, "wrap": False},
    {"key": "scheduled_date", "title": "Дата", "width": 14, "locked": False, "hidden": False, "wrap": False,
     "hint": "Формат: 11.09.2026. Можно оставить пустой — дата рассчитается по расписанию плана."},
    {"key": "title", "title": "Название урока", "width": 28, "locked": False, "hidden": False, "wrap": True,
     "hint": "Можно оставить пустым, если заполнена тема."},
    {"key": "topic", "title": "Тема", "width": 28, "locked": False, "hidden": False, "wrap": True,
     "hint": "Например: Системы счисления"},
    {"key": "subtopic", "title": "Подтема", "width": 28, "locked": False, "hidden": False, "wrap": True,
     "hint": "Например: Перевод чисел между системами счисления"},
    {"key": "task_number", "title": "№ задания", "width": 14, "locked": False, "hidden": False, "wrap": False,
     "hint": "Номер задания ЕГЭ/ОГЭ, например 5, 8, 14 или 1-5."},
    {"key": "goal", "title": "Цель", "width": 36, "locked": False, "hidden": False, "wrap": True,
     "hint": "Что ученик должен уметь после занятия."},
    {"key": "description", "title": "План урока", "width": 40, "locked": False, "hidden": False, "wrap": True,
     "hint": "Краткий последовательный план занятия."},
    {"key": "homework_description", "title": "Домашнее задание", "width": 36, "locked": False, "hidden": False, "wrap": True,
     "hint": "Текстовое описание ДЗ. Файлы и задания добавляются потом в карточке урока."},
    {"key": "teacher_comment", "title": "Комментарий", "width": 28, "locked": False, "hidden": False, "wrap": True,
     "hint": "Заметка учителя. Не видна ученику как материал."},
    {"key": "date_source", "title": "Статус даты", "width": 14, "locked": True, "hidden": True, "wrap": False},
    {"key": "order", "title": "Порядок", "width": 10, "locked": True, "hidden": True, "wrap": False},
)

HEADER_ALIASES = {
    "№": "number",
    "ID урока": "item_id",
    "Дата": "scheduled_date",
    "Дата занятия": "scheduled_date",
    "Название урока": "title",
    "Тема": "topic",
    "Подтема": "subtopic",
    "№ задания": "task_number",
    "Цель": "goal",
    "План урока": "description",
    "Домашнее задание": "homework_description",
    "Комментарий": "teacher_comment",
    "Статус даты": "date_source",
    "Порядок": "order",
}
REQUIRED_HEADERS = ("Название урока", "Тема", "Дата")
TITLE_BY_KEY = {col["key"]: col["title"] for col in COLUMNS}
USER_KEYS = {col["key"] for col in COLUMNS if not col["locked"]}

SETTINGS_FIELDS = (
    {
        "key": "subject",
        "label": "Предмет",
        "hint": "Выберите предмет из списка. При импорте уроков это поле обновляет карточку плана.",
        "imported": True,
    },
    {
        "key": "direction",
        "label": "Уровень",
        "hint": "ЕГЭ, ОГЭ и другие уровни плана. При импорте обновляет карточку плана.",
        "imported": True,
    },
    {
        "key": "grade",
        "label": "Класс",
        "hint": "Например: 9 или 10–11. При импорте обновляет карточку плана.",
        "imported": True,
    },
    {
        "key": "start_date",
        "label": "Дата начала",
        "hint": "Первая дата расписания. Нужна, если в таблице уроков дата не указана.",
        "imported": True,
    },
    {
        "key": "interval",
        "label": "Расписание",
        "hint": "Как часто идут занятия. По этому расписанию считаются автоматические даты и сдвиг уроков.",
        "imported": True,
    },
)


class PlanExcelError(Exception):
    def __init__(self, message, *, code="excel_error", http_status=status.HTTP_400_BAD_REQUEST):
        super().__init__(message)
        self.message = message
        self.code = code
        self.http_status = http_status


def normalize_import_mode(value: Any) -> str:
    raw = str(value or "").strip().lower()
    if raw in LEGACY_MODE_MAP:
        return LEGACY_MODE_MAP[raw]
    if raw in IMPORT_MODES:
        return raw
    raise PlanExcelError("Неизвестный режим импорта.")


def normalize_topic_key(value: str) -> str:
    return " ".join(str(value or "").split()).casefold()


def excel_safe_text(value: Any) -> str:
    text = "" if value is None else str(value)
    if text[:1] in {"=", "+", "-", "@"}:
        return f"'{text}"
    return text


def format_ru_date(value: Any) -> str:
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    else:
        parsed = parse_plan_date(value)
    if not parsed:
        return ""
    return parsed.strftime("%d.%m.%Y")


def _clip(value: str, limit: int) -> tuple[str, bool]:
    text = str(value or "")
    if len(text) <= limit:
        return text, False
    return text[:limit], True


def _lessons_word(count: int) -> str:
    n = abs(int(count or 0))
    if n % 10 == 1 and n % 100 != 11:
        return "урок"
    if 2 <= n % 10 <= 4 and not (12 <= n % 100 <= 14):
        return "урока"
    return "уроков"


def parse_excel_date(value: Any) -> tuple[date | None, str | None]:
    if value is None or value == "":
        return None, None
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    if isinstance(value, bool):
        return None, "некорректная дата"
    if isinstance(value, (int, float)):
        try:
            parsed = from_excel(value)
        except Exception:
            return None, "некорректная дата"
        if isinstance(parsed, datetime):
            return parsed.date(), None
        if isinstance(parsed, date):
            return parsed, None
        return None, "некорректная дата"
    raw = str(value).strip()
    if not raw:
        return None, None
    if raw.startswith("="):
        return None, "формулы в дате не поддерживаются"
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y", "%d/%m/%y"):
        try:
            return datetime.strptime(raw, fmt).date(), None
        except ValueError:
            continue
    iso = parse_plan_date(raw)
    if iso:
        return iso, None
    return None, f"некорректная дата «{raw}»"


def template_filename(*, title="", subject="", direction="", filled=False) -> str:
    subject_label = get_plan_subject_label(subject) if subject else ""
    level_label = get_plan_level_label(direction) if direction else ""
    if filled and title:
        base = f"План_уроков_{title}"
    elif subject_label and level_label:
        base = f"План_уроков_{subject_label}_{level_label}"
    elif subject_label:
        base = f"План_уроков_{subject_label}"
    else:
        base = "Шаблон_плана_уроков"
    safe = sanitize_filename(re.sub(r"\s+", "_", str(base).strip()))
    if not safe.lower().endswith(".xlsx"):
        safe = f"{safe}.xlsx"
    return safe


def xlsx_response(content: bytes, filename: str) -> HttpResponse:
    response = HttpResponse(
        content,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = content_disposition(filename, inline=False)
    return response


def build_plan_workbook(settings: dict | None = None, items: list | None = None) -> bytes:
    settings = settings or {}
    items = list(items or [])
    wb = Workbook()
    lessons = wb.active
    lessons.title = SHEET_LESSONS
    _write_lessons_sheet(lessons, items)
    _write_instructions_sheet(wb.create_sheet(SHEET_INSTRUCTIONS))
    settings_ws = wb.create_sheet(SHEET_SETTINGS)
    meta = wb.create_sheet(SHEET_META)
    _write_meta_sheet(meta, settings)
    _write_settings_sheet(settings_ws, settings, meta)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_lessons_sheet(ws, items: list) -> None:
    last_col = get_column_letter(len(COLUMNS))
    blank_count = max(len(items) + TEMPLATE_SPARE_ROWS, TEMPLATE_SPARE_ROWS, 2)
    last_row = 1 + blank_count
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24
    ws.sheet_view.showGridLines = False
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    for index, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, index, col["title"])
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center")
        cell.border = THIN
        cell.protection = Protection(locked=True)
        if col.get("hint"):
            cell.comment = Comment(col["hint"], "Цифровой поток")
        ws.column_dimensions[get_column_letter(index)].width = col["width"]
        ws.column_dimensions[get_column_letter(index)].hidden = bool(col["hidden"])

    for row_idx in range(2, last_row + 1):
        item = items[row_idx - 2] if row_idx - 2 < len(items) else None
        ws.row_dimensions[row_idx].height = 32 if item else 22
        banded = BAND_FILL if row_idx % 2 == 0 else WHITE_FILL
        for col_idx, col in enumerate(COLUMNS, start=1):
            cell = ws.cell(row_idx, col_idx)
            cell.font = CELL_FONT
            cell.border = THIN
            cell.fill = banded
            cell.protection = Protection(locked=col["locked"])
            cell.alignment = WRAP if col["wrap"] else CENTER
            if col["key"] == "scheduled_date":
                cell.number_format = "DD.MM.YYYY"
            elif col["key"] not in {"number", "order"}:
                cell.number_format = "@"
            if col["key"] == "number":
                cell.value = "=ROW()-1"
                cell.protection = Protection(locked=True)
            elif item:
                cell.value = _export_cell(col["key"], item, row_idx - 1)

    date_letters = get_column_letter(next(i for i, col in enumerate(COLUMNS, start=1) if col["key"] == "scheduled_date"))
    dv = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2000,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=True,
        showInputMessage=True,
        promptTitle="Дата",
        prompt="Формат: ДД.ММ.ГГГГ. Можно оставить пустой.",
        showErrorMessage=False,
    )
    dv.add(f"{date_letters}2:{date_letters}{last_row}")
    ws.add_data_validation(dv)

    table = Table(displayName="PlanLessons", ref=f"A1:{last_col}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight1",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    ws.protection.sheet = True
    ws.protection.enable()
    ws.protection.insertRows = True
    ws.protection.deleteRows = True
    ws.protection.insertColumns = False
    ws.protection.deleteColumns = False
    ws.protection.autoFilter = True
    ws.protection.sort = True
    ws.protection.formatCells = True
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True


def _export_cell(key: str, item: dict, order: int):
    if key == "item_id":
        return item.get("id") or ""
    if key == "order":
        return item.get("order") or order
    if key == "date_source":
        return item.get("date_source") or ""
    if key == "scheduled_date":
        value = item.get("scheduled_date")
        parsed = parse_plan_date(value) if not isinstance(value, date) else value
        return parsed
    raw = item.get(key) or ""
    if key in USER_KEYS:
        return excel_safe_text(raw)
    return raw


def _write_instructions_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 36
    ws.column_dimensions["E"].width = 36
    ws.merge_cells("A1:E1")
    ws["A1"] = "План занятий в Excel"
    ws["A1"].font = TITLE_FONT
    ws.row_dimensions[1].height = 28

    row = 3
    row = _instr_section(ws, row, "1. Для чего нужен файл")
    row = _instr_box(
        ws,
        row,
        "Шаблон позволяет подготовить или изменить план занятий в Excel и загрузить его в Цифровой поток. "
        "Вы заполняете привычную таблицу, а на платформе выбираете, что сделать: заменить план целиком, "
        "добавить уроки или обновить конкретные занятия по датам.",
    )
    row += 1
    row = _instr_section(ws, row, "2. Как заполнить таблицу")
    row = _instr_para(ws, row, "Один урок — одна строка на листе «Уроки». Полностью пустые строки игнорируются.")
    row = _instr_table(
        ws,
        row,
        ("Столбец", "Обязательный?", "Что означает", "Пример"),
        (
            ("Дата", "Зависит от режима", "Формат 11.09.2026. Если оставить пустой, дата может быть рассчитана автоматически там, где это поддерживает расписание плана.", "11.09.2026"),
            ("Название урока", "Нужно название или тема", "Можно оставить пустым, если заполнена тема.", "Урок 3"),
            ("Тема", "Нужно название или тема", "Например: Системы счисления. Регистр и лишние пробелы не важны при точном сопоставлении. Похожие названия автоматически не подменяются.", "Системы счисления"),
            ("Подтема", "Нет", "Уточнение темы занятия.", "Перевод чисел"),
            ("№ задания", "Нет", "Номер задания экзамена, если нужен.", "5 или 1-5"),
            ("Цель", "Нет", "Что ученик должен уметь после занятия.", "Переводить числа между системами счисления"),
            ("План урока", "Нет", "Краткий ход занятия.", "Повторение, разбор, практика"),
            ("Домашнее задание", "Нет", "Текстовое описание. Файлы добавляются уже на платформе.", "№ 5, 8"),
            ("Комментарий", "Нет", "Заметка учителя, не видна ученику как материал.", "Повторить перевод в 16-ричную"),
        ),
    )
    row += 1
    row = _instr_section(ws, row, "3. Что считается пустой строкой")
    row = _instr_box(ws, row, "Полностью пустые строки игнорируются. Если в строке есть хотя бы дата, название, тема или другой текст урока — строка считается занятой.")
    row += 1
    row = _instr_section(ws, row, "4. Как работают даты")
    row = _instr_box(
        ws,
        row,
        "Пишите дату как ДД.ММ.ГГГГ, например 11.09.2026. Дата из ячейки Excel Date тоже подходит: 11.09.2026 всегда означает календарный день 2026-09-11, часовой пояс его не сдвигает.\n\n"
        "Если дата указана в файле, она считается ручной (manual): платформа сохранит именно этот день.\n"
        "Если дата пустая, она считается автоматической: её рассчитает расписание плана (раз в неделю, 2 раза в неделю и т. д.).",
        height=72,
    )
    row += 1
    row = _instr_section(ws, row, "5. Три режима импорта")
    row = _instr_para(ws, row, "Режим выбирается на платформе при загрузке. В Excel не нужно писать CREATE/UPDATE или ID урока.")
    row = _instr_para(ws, row, "Обновить план целиком. Текущий список уроков полностью заменяется содержимым файла. Если в плане было 20 уроков, а в файле 8 — после импорта останется 8. Порядок строк Excel становится порядком уроков. Старые уроки, которых нет в файле, не сохраняются. Строки не сопоставляются по названию, теме или номеру.")
    row = _instr_example(
        ws,
        row,
        "Было в плане",
        (("11.09", "Урок 1"), ("16.09", "Урок 2"), ("18.09", "Урок 3")),
        "Excel",
        (("11.09", "Новый A"), ("18.09", "Новый B")),
        "После импорта",
        (("11.09", "Новый A"), ("18.09", "Новый B")),
    )
    row = _instr_para(ws, row, "Добавить уроки. Уроки из Excel вставляются на указанные даты. Последующие занятия сдвигаются дальше по плану согласно расписанию, а не простым «плюс несколько дней». Ручные даты остаются ручными. ID урока в этом режиме никогда не обновляет существующий урок — только добавляет новый.")
    row = _instr_example(
        ws,
        row,
        "Было в плане",
        (("11.09", "Урок A"), ("16.09", "Урок B"), ("18.09", "Урок C")),
        "Excel",
        (("16.09", "Новый урок"),),
        "После импорта",
        (("11.09", "Урок A"), ("16.09", "Новый урок"), ("далее", "Урок B"), ("далее", "Урок C")),
    )
    row = _instr_para(ws, row, "Заменить уроки по датам. Это не полная замена плана. Для каждой строки Excel находится урок, который уже стоит на этой дате, и заменяется только его содержимое. Количество уроков не меняется. Если на дату урока нет, строка не станет новым уроком — это ошибка. Если на одну дату несколько уроков или в файле две строки на одну дату — тоже ошибка.")
    row = _instr_example(
        ws,
        row,
        "Было в плане",
        (("11.09", "Системы счисления"), ("16.09", "Логика"), ("18.09", "Кодирование")),
        "Excel",
        (("16.09", "Алгебра логики"),),
        "После импорта",
        (("11.09", "Системы счисления"), ("16.09", "Алгебра логики"), ("18.09", "Кодирование")),
    )
    row += 1
    row = _instr_section(ws, row, "6. Что произойдёт при ошибке")
    row = _instr_box(
        ws,
        row,
        "Сначала показывается предпросмотр. Если в нём есть критическая ошибка, импорт не применяется частично: ничего не удаляется, ничего не создаётся, текущий план остаётся без изменений. Исправьте файл и загрузите его снова.",
        fill=WARN_FILL,
    )
    row += 1
    row = _instr_section(ws, row, "7. Чего нельзя делать")
    row = _instr_para(
        ws,
        row,
        "• не переименовывать обязательные заголовки (Дата, Название урока, Тема);\n"
        "• не удалять обязательные столбцы;\n"
        "• не использовать .xls / .xlsm — нужен только .xlsx;\n"
        "• не добавлять VBA и макросы;\n"
        "• не редактировать скрытые системные поля без необходимости.",
        height=86,
    )
    row += 1
    row = _instr_section(ws, row, "8. Пошаговая загрузка")
    row = _instr_para(
        ws,
        row,
        "1. Заполните лист Уроки.\n"
        "2. При необходимости проверьте Настройки.\n"
        "3. Сохраните файл в формате .xlsx.\n"
        "4. Откройте план на платформе.\n"
        "5. Excel → Импортировать.\n"
        "6. Выберите файл.\n"
        "7. Выберите режим.\n"
        "8. Проверьте предпросмотр.\n"
        "9. Нажмите кнопку применения.",
        height=150,
    )
    row += 1
    row = _instr_section(ws, row, "9. Примеры")
    row = _instr_para(ws, row, "Пример А. Новый план на четверть: заполните 8 строк и выберите «Обновить план целиком». Старый список из 20 уроков будет заменён этими 8.")
    row = _instr_para(ws, row, "Пример Б. Нужно вставить контрольную на 16.09: одна строка с датой 16.09.2026 и режимом «Добавить уроки». Урок встанет на 16.09, а следующие сдвинутся по расписанию.")
    row = _instr_para(ws, row, "Пример В. На 16.09 уже стоит «Логика», нужно заменить тему на «Алгебра логики»: одна строка с этой датой и режимом «Заменить уроки по датам». Остальные уроки не изменятся.")

    ws.protection.sheet = True
    ws.protection.enable()
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True


def _instr_section(ws, row: int, title: str) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    cell = ws.cell(row, 1, title)
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL
    cell.alignment = CENTER
    for col in range(1, 6):
        ws.cell(row, col).fill = SECTION_FILL
        ws.cell(row, col).border = THIN
    ws.row_dimensions[row].height = 22
    return row + 1


def _instr_para(ws, row: int, text: str, *, height=48) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    cell = ws.cell(row, 1, text)
    cell.font = CELL_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = height
    return row + 1


def _instr_box(ws, row: int, text: str, *, height=56, fill=None) -> int:
    fill = fill or BOX_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    cell = ws.cell(row, 1, text)
    cell.font = CELL_FONT
    cell.fill = fill
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    for col in range(1, 6):
        ws.cell(row, col).fill = fill
        ws.cell(row, col).border = THIN
    ws.row_dimensions[row].height = height
    return row + 1


def _instr_table(ws, row: int, headers: tuple, rows: tuple) -> int:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row, col, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN
    ws.row_dimensions[row].height = 20
    row += 1
    for index, values in enumerate(rows):
        fill = BAND_FILL if index % 2 == 0 else WHITE_FILL
        height = 48 if any(len(str(value)) > 60 for value in values) else 32
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            cell.font = CELL_FONT
            cell.fill = fill
            cell.alignment = WRAP
            cell.border = THIN
        ws.row_dimensions[row].height = height
        row += 1
    return row


def _instr_example(ws, row: int, left_title: str, left_rows: tuple, mid_title: str, mid_rows: tuple, right_title: str, right_rows: tuple) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=3)
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
    for col, title in ((1, left_title), (3, mid_title), (4, right_title)):
        cell = ws.cell(row, col, title)
        cell.font = LABEL_FONT
        cell.fill = BOX_FILL
        cell.alignment = CENTER
    ws.cell(row, 2).fill = BOX_FILL
    ws.cell(row, 5).fill = BOX_FILL
    row += 1
    count = max(len(left_rows), len(mid_rows), len(right_rows))
    for index in range(count):
        left = left_rows[index] if index < len(left_rows) else ("", "")
        mid = mid_rows[index] if index < len(mid_rows) else ("", "")
        right = right_rows[index] if index < len(right_rows) else ("", "")
        ws.cell(row, 1, left[0]).font = CELL_FONT
        ws.cell(row, 2, left[1]).font = CELL_FONT
        ws.cell(row, 3, f"{mid[0]}  {mid[1]}".strip()).font = CELL_FONT
        ws.cell(row, 4, right[0]).font = CELL_FONT
        ws.cell(row, 5, right[1]).font = CELL_FONT
        for col in range(1, 6):
            ws.cell(row, col).border = THIN
            ws.cell(row, col).alignment = CENTER
        row += 1
    return row + 1


def _write_settings_sheet(ws, settings: dict, meta) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 72
    ws.merge_cells("A1:C1")
    ws["A1"] = "НАСТРОЙКИ ПЛАНА"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A2:C2")
    ws["A2"] = "Редактируйте значения в столбце B. Названия полей и подсказки защищены от случайного изменения."
    ws["A2"].font = HINT_FONT
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 28

    subject_label = get_plan_subject_label(settings.get("subject") or "") or settings.get("subject") or ""
    level_label = get_plan_level_label(settings.get("direction") or "") or settings.get("direction") or ""
    interval_label = INTERVAL_LABELS.get(
        normalize_interval(settings.get("interval")),
        settings.get("interval") or "",
    )
    values = {
        "subject": subject_label,
        "direction": level_label,
        "grade": settings.get("grade") or "",
        "start_date": parse_plan_date(settings.get("start_date")) or settings.get("start_date") or "",
        "interval": interval_label,
    }
    start_row = 4
    for offset, field in enumerate(SETTINGS_FIELDS):
        row = start_row + offset
        label_cell = ws.cell(row, 1, field["label"])
        label_cell.font = LABEL_FONT
        label_cell.fill = BOX_FILL
        label_cell.alignment = CENTER
        label_cell.border = THIN
        label_cell.protection = Protection(locked=True)
        value_cell = ws.cell(row, 2, values[field["key"]])
        value_cell.font = CELL_FONT
        value_cell.alignment = CENTER
        value_cell.border = THIN
        value_cell.protection = Protection(locked=False)
        value_cell.fill = WHITE_FILL
        if field["key"] == "start_date":
            value_cell.number_format = "DD.MM.YYYY"
        hint_cell = ws.cell(row, 3, field["hint"])
        hint_cell.font = HINT_FONT
        hint_cell.alignment = Alignment(wrap_text=True, vertical="center")
        hint_cell.protection = Protection(locked=True)
        ws.row_dimensions[row].height = 28

    _add_settings_validation(ws, meta, start_row)

    note_row = start_row + len(SETTINGS_FIELDS) + 1
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 1, end_column=3)
    note = ws.cell(
        note_row,
        1,
        "Предмет, уровень и класс обновляют карточку плана. Дата начала и расписание используются, "
        "чтобы рассчитать пустые даты уроков и сдвинуть последующие занятия в режиме «Добавить уроки».",
    )
    note.font = HINT_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")
    note.fill = BOX_FILL
    for col in range(1, 4):
        ws.cell(note_row, col).fill = BOX_FILL
        ws.cell(note_row + 1, col).fill = BOX_FILL
        ws.cell(note_row, col).border = THIN
        ws.cell(note_row + 1, col).border = THIN
    ws.row_dimensions[note_row].height = 22
    ws.row_dimensions[note_row + 1].height = 22

    ws.protection.sheet = True
    ws.protection.enable()
    ws.protection.insertRows = False
    ws.protection.deleteRows = False
    ws.protection.insertColumns = False
    ws.protection.deleteColumns = False
    ws.protection.formatCells = True
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True


def _add_settings_validation(ws, meta, start_row: int) -> None:
    subject_labels = [item["label"] for item in get_plan_subject_options()]
    level_labels = [item["label"] for item in get_plan_level_options()]
    interval_labels = list(INTERVAL_LABELS.values())
    _write_meta_list(meta, 4, "subjects", subject_labels)
    _write_meta_list(meta, 5, "levels", level_labels)
    _write_meta_list(meta, 6, "grades", list(GRADE_CHOICES))
    _write_meta_list(meta, 7, "intervals", interval_labels)

    def list_dv(col_letter: str, last_row: int, title: str, prompt: str):
        formula = f"'{SHEET_META}'!${col_letter}$2:${col_letter}${last_row}"
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=True,
            showDropDown=False,
            showErrorMessage=False,
            showInputMessage=True,
            promptTitle=title,
            prompt=prompt,
        )
        return dv

    subject_dv = list_dv("D", 1 + max(len(subject_labels), 1), "Предмет", "Выберите предмет плана.")
    subject_dv.add(f"B{start_row}")
    ws.add_data_validation(subject_dv)
    level_dv = list_dv("E", 1 + max(len(level_labels), 1), "Уровень", "Выберите уровень плана.")
    level_dv.add(f"B{start_row + 1}")
    ws.add_data_validation(level_dv)
    grade_dv = list_dv("F", 1 + len(GRADE_CHOICES), "Класс", "Выберите класс или введите свой вариант.")
    grade_dv.add(f"B{start_row + 2}")
    ws.add_data_validation(grade_dv)
    date_dv = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2000,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=True,
        showErrorMessage=False,
        showInputMessage=True,
        promptTitle="Дата начала",
        prompt="Формат: ДД.ММ.ГГГГ.",
    )
    date_dv.add(f"B{start_row + 3}")
    ws.add_data_validation(date_dv)
    interval_dv = list_dv("G", 1 + len(interval_labels), "Расписание", "Как часто идут занятия.")
    interval_dv.add(f"B{start_row + 4}")
    ws.add_data_validation(interval_dv)


def _write_meta_list(ws, column: int, name: str, values: list[str]) -> None:
    ws.cell(1, column, name)
    for index, value in enumerate(values, start=2):
        ws.cell(index, column, value)


def _write_meta_sheet(ws, settings: dict) -> None:
    ws.sheet_state = "hidden"
    ws["A1"] = "template_version"
    ws["B1"] = TEMPLATE_VERSION
    ws["A2"] = "plan_id"
    ws["B2"] = settings.get("plan_id") or ""
    ws["A3"] = "subject"
    ws["B3"] = settings.get("subject") or ""
    ws["A4"] = "direction"
    ws["B4"] = settings.get("direction") or ""
    ws["A5"] = "interval"
    ws["B5"] = settings.get("interval") or ""
    ws["A6"] = "start_date"
    ws["B6"] = settings.get("start_date") or ""
    ws["A7"] = "grade"
    ws["B7"] = settings.get("grade") or ""


def read_uploaded_bytes(uploaded) -> bytes:
    if uploaded is None:
        raise PlanExcelError("Выберите файл Excel.")
    name = str(getattr(uploaded, "name", "") or "")
    if name and not name.lower().endswith(".xlsx"):
        raise PlanExcelError("Нужен файл в формате .xlsx. CSV, XLS и XLSM не подходят.")
    size = getattr(uploaded, "size", None)
    if size is not None and size > MAX_FILE_BYTES:
        raise PlanExcelError("Файл слишком большой. Загрузите Excel размером до 4 МБ.")
    data = uploaded.read()
    if not data:
        raise PlanExcelError("Файл пустой.")
    if len(data) > MAX_FILE_BYTES:
        raise PlanExcelError("Файл слишком большой. Загрузите Excel размером до 4 МБ.")
    if data[:2] != b"PK":
        raise PlanExcelError("Файл не похож на Excel-шаблон. Сохраните его как .xlsx.")
    return data


def parse_plan_workbook(
    data: bytes,
    *,
    existing_items: list | None = None,
    start_date: Any = None,
    interval: str = "weekly",
    extra_topics: list[str] | None = None,
    mode: str = MODE_REPLACE_ALL,
) -> dict:
    parsed = _parse_workbook_rows(data, extra_topics=extra_topics)
    return build_excel_import_plan(
        parsed,
        existing_items=existing_items,
        start_date=start_date,
        interval=interval,
        mode=mode,
    )


def _parse_workbook_rows(data: bytes, *, extra_topics: list[str] | None = None) -> dict:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=False, keep_vba=False)
    except Exception as exc:
        logger.info("plan excel open failed: %s", exc)
        raise PlanExcelError("Не удалось прочитать файл. Убедитесь, что это .xlsx без макросов.") from exc
    if getattr(wb, "vba_archive", None):
        raise PlanExcelError("Файлы с макросами не поддерживаются. Сохраните шаблон как .xlsx.")

    version = _read_template_version(wb)
    if version not in SUPPORTED_TEMPLATE_VERSIONS:
        raise PlanExcelError(
            f"Этот шаблон версии {version} пока не поддерживается. Скачайте актуальный шаблон."
        )
    file_settings = _read_file_settings(wb)
    ws, header_map, extra_headers = _find_lessons_sheet(wb)
    warnings = []
    if extra_headers:
        extras = ", ".join(extra_headers[:6])
        warnings.append(f"Лишние столбцы проигнорированы: {extras}.")

    topic_canon = {}
    for topic in extra_topics or []:
        key = normalize_topic_key(topic)
        if key and key not in topic_canon:
            topic_canon[key] = str(topic).strip()

    rows = []
    skipped_empty = 0
    pending_empty = 0
    seen_lesson = False
    for excel_row in range(2, (ws.max_row or 1) + 1):
        raw = {key: ws.cell(excel_row, col_idx).value for key, col_idx in header_map.items()}
        if _row_is_empty(raw):
            if seen_lesson:
                pending_empty += 1
            continue
        if pending_empty:
            skipped_empty += pending_empty
            pending_empty = 0
        seen_lesson = True
        parsed = _parse_lesson_row(raw, excel_row=excel_row, topic_canon=topic_canon)
        rows.append(parsed)
        if len(rows) > MAX_LESSONS:
            raise PlanExcelError(
                f"В файле больше {MAX_LESSONS} уроков. За один раз можно импортировать не более {MAX_LESSONS}."
            )

    if not rows:
        raise PlanExcelError("В файле нет уроков для импорта. Заполните хотя бы одну строку.")

    return {
        "template_version": version or TEMPLATE_VERSION,
        "warnings": warnings,
        "rows": rows,
        "skipped_empty": skipped_empty,
        "settings": file_settings,
    }


def build_excel_import_plan(
    parsed: dict,
    *,
    existing_items: list | None = None,
    start_date: Any = None,
    interval: str = "weekly",
    mode: str = MODE_REPLACE_ALL,
) -> dict:
    mode = normalize_import_mode(mode)
    file_settings = dict(parsed.get("settings") or {})
    start = (
        parse_plan_date(start_date)
        or parse_plan_date(file_settings.get("start_date"))
        or _first_explicit_date(parsed.get("rows") or [])
    )
    interval = normalize_interval(interval or file_settings.get("interval") or "weekly")
    existing_entries = _existing_entries(existing_items or [], interval)
    topic_canon = {}
    for entry in existing_entries:
        key = normalize_topic_key(entry.get("topic") or "")
        if key and key not in topic_canon:
            topic_canon[key] = str(entry.get("topic") or "").strip()
    rows = [_with_canonical_topic(row, topic_canon) for row in (parsed.get("rows") or [])]

    if mode == MODE_REPLACE_ALL:
        plan = _resolve_replace_all(rows, existing_entries, start, interval)
    elif mode == MODE_INSERT:
        plan = _resolve_insert(rows, existing_entries, start, interval)
    else:
        plan = _resolve_replace_dates(rows, existing_entries)

    warnings = list(parsed.get("warnings") or []) + list(plan.get("warnings") or [])
    errors = sum(1 for row in plan["rows"] if row["status"] == "error")
    warn_rows = sum(1 for row in plan["rows"] if row["status"] == "warning")
    can_import = errors == 0 and not plan.get("blocking_errors")
    auto_dates = sum(1 for row in plan["rows"] if row["item"].get("date_source") == "automatic" and row["item"].get("scheduled_date"))
    manual_dates = sum(1 for row in plan["rows"] if row["item"].get("date_source") == "manual")
    summary = {
        "found": len(rows),
        "file_count": len(rows),
        "current_count": len(existing_entries),
        "after_count": plan.get("after_count", len(existing_entries)),
        "added": plan.get("added", 0),
        "replaced": plan.get("replaced", 0),
        "skip": parsed.get("skipped_empty") or 0,
        "errors": errors + len(plan.get("blocking_errors") or []),
        "warnings": warn_rows + len(warnings),
        "auto_dates": auto_dates,
        "manual_dates": manual_dates,
        "subsequent_shift": bool(plan.get("shifts")),
    }
    return {
        "ok": True,
        "mode": mode,
        "template_version": parsed.get("template_version") or TEMPLATE_VERSION,
        "warnings": warnings,
        "blocking_errors": plan.get("blocking_errors") or [],
        "confirmation": plan.get("confirmation") or "",
        "summary": summary,
        "rows": plan["rows"],
        "shifts": plan.get("shifts") or [],
        "operations": plan.get("operations") or [],
        "settings": {**file_settings, "start_date": start.isoformat() if start else file_settings.get("start_date") or "", "interval": interval},
        "can_import": can_import,
        "start_date": start.isoformat() if start else "",
        "interval": interval,
    }


def apply_plan_excel_import(plan: LessonPlan, preview: dict, *, mode: str, teacher) -> dict:
    mode = normalize_import_mode(mode or preview.get("mode"))
    from rest_framework.exceptions import ValidationError as DRFValidationError
    from .serializers import LessonPlanItemEditorSerializer

    if preview.get("mode") and preview["mode"] != mode:
        raise PlanExcelError("Предпросмотр построен для другого режима. Загрузите файл снова.")
    if not preview.get("can_import"):
        first = next((row for row in preview.get("rows") or [] if row.get("status") == "error"), None)
        message = ""
        if first and first.get("messages"):
            message = str(first["messages"][0])
        elif preview.get("blocking_errors"):
            message = str(preview["blocking_errors"][0])
        raise PlanExcelError(message or "Сначала исправьте ошибки в файле.")

    operations = list(preview.get("operations") or [])
    created = 0
    updated = 0
    deleted = 0
    item_dates = []

    with transaction.atomic():
        existing = list(plan.items.select_related("scheduled_event").prefetch_related("schedule_events_linked").order_by("order", "id"))
        existing_by_id = {item.pk: item for item in existing}

        if mode == MODE_REPLACE_ALL:
            blocked = [item for item in existing if _item_in_use(item)]
            if blocked:
                titles = ", ".join(item.title for item in blocked[:3])
                raise PlanExcelError(f"Нельзя обновить план целиком: занятие «{titles}» уже стоит в расписании.")
            for item in existing:
                item.delete()
                deleted += 1
            existing_by_id = {}

        for operation in operations:
            op = operation.get("op")
            payload = _serializer_payload(operation.get("item") or {}, operation.get("order") or 1)
            target_id = operation.get("id")
            instance = existing_by_id.get(int(target_id)) if target_id else None
            if instance is None and target_id:
                raise PlanExcelError("Нельзя изменить урок, который не относится к текущему плану.")
            try:
                if op == "create" or instance is None:
                    serializer = LessonPlanItemEditorSerializer(data=payload, context={"teacher": teacher})
                    serializer.is_valid(raise_exception=True)
                    item = serializer.save(plan=plan, status=_status_for_date(payload.get("scheduled_date")))
                    created += 1
                elif op == "update":
                    serializer = LessonPlanItemEditorSerializer(
                        instance,
                        data=payload,
                        partial=True,
                        context={"teacher": teacher},
                    )
                    serializer.is_valid(raise_exception=True)
                    item = serializer.save(status=_status_for_date(payload.get("scheduled_date")) or instance.status)
                    updated += 1
                else:
                    continue
            except DRFValidationError as exc:
                raise PlanExcelError("Не удалось сохранить урок. Проверьте названия и длины полей.") from exc
            item_dates.append({
                "id": item.pk,
                "date_source": (operation.get("item") or {}).get("date_source") or "",
            })

        _apply_plan_settings(plan, preview.get("settings") or {})
        plan.lessons_count = plan.items.count()
        plan.save(update_fields=["lessons_count", "updated_at", "subject", "direction", "grade"])

    return {
        "created": created,
        "updated": updated,
        "deleted": deleted,
        "mode": mode,
        "auto_dates": preview["summary"].get("auto_dates") or 0,
        "manual_dates": preview["summary"].get("manual_dates") or 0,
        "skipped": preview["summary"].get("skip") or 0,
        "item_dates": item_dates,
    }


def _apply_plan_settings(plan: LessonPlan, settings: dict) -> None:
    subject = str(settings.get("subject") or "").strip()
    direction = str(settings.get("direction") or "").strip()
    grade = str(settings.get("grade") or "").strip()
    if subject:
        plan.subject = subject[:20]
    if direction:
        plan.direction = direction[:20]
    if grade:
        plan.grade = grade[:32]


def _status_for_date(scheduled):
    return PlanItemStatus.PLANNED if scheduled else PlanItemStatus.NOT_STARTED


def _item_in_use(item) -> bool:
    if isinstance(item, dict):
        return bool(item.get("in_use"))
    if getattr(item, "scheduled_event_id", None):
        return True
    linked = item.schedule_events_linked.exclude(status=ScheduleEvent.Status.CANCELLED)
    return linked.exists()


def _serializer_payload(item: dict, order: int) -> dict:
    scheduled = item.get("scheduled_date")
    if hasattr(scheduled, "isoformat"):
        scheduled = scheduled.isoformat()
    return {
        "order": order,
        "title": item.get("title") or "",
        "topic": item.get("topic") or "",
        "subtopic": item.get("subtopic") or "",
        "task_number": item.get("task_number") or "",
        "goal": item.get("goal") or "",
        "description": item.get("description") or "",
        "homework_description": item.get("homework_description") or "",
        "teacher_comment": item.get("teacher_comment") or "",
        "scheduled_date": scheduled or None,
    }


def _read_template_version(wb) -> str:
    meta = _read_meta_settings(wb)
    return str(meta.get("template_version") or "").strip()


def _read_meta_settings(wb) -> dict:
    settings = {}
    if SHEET_META not in wb.sheetnames:
        return settings
    meta = wb[SHEET_META]
    for row in meta.iter_rows(min_row=1, max_row=8, max_col=2, values_only=True):
        key = str(row[0] or "").strip()
        if key:
            settings[key] = row[1]
    return settings


def _read_file_settings(wb) -> dict:
    settings = _read_meta_settings(wb)
    parsed = {
        "subject": normalize_plan_subject_id(settings.get("subject") or "") or "",
        "direction": normalize_plan_level_id(settings.get("direction") or "") or "",
        "grade": str(settings.get("grade") or "").strip(),
        "start_date": settings.get("start_date") or "",
        "interval": settings.get("interval") or "",
    }
    if SHEET_SETTINGS in wb.sheetnames:
        ws = wb[SHEET_SETTINGS]
        by_label = {field["label"]: field["key"] for field in SETTINGS_FIELDS}
        for row in ws.iter_rows(min_row=1, max_row=20, max_col=2, values_only=True):
            label = str(row[0] or "").strip()
            key = by_label.get(label)
            if not key:
                continue
            parsed[key] = row[1]
    return _normalize_settings_values(parsed)


def _normalize_settings_values(raw: dict) -> dict:
    subject = _match_option(raw.get("subject"), get_plan_subject_options(), normalize_plan_subject_id)
    direction = _match_option(raw.get("direction"), get_plan_level_options(), normalize_plan_level_id)
    interval_raw = raw.get("interval")
    interval = ""
    if interval_raw not in (None, ""):
        label_map = {label.casefold(): key for key, label in INTERVAL_LABELS.items()}
        text = str(interval_raw).strip()
        interval = label_map.get(text.casefold()) or normalize_interval(text)
    start, _err = parse_excel_date(raw.get("start_date"))
    grade = str(raw.get("grade") or "").strip()
    return {
        "subject": subject,
        "direction": direction,
        "grade": grade,
        "start_date": start.isoformat() if start else "",
        "interval": interval,
    }


def _match_option(value, options, normalize) -> str:
    if value in (None, ""):
        return ""
    raw = str(value).strip()
    nid = normalize(raw)
    for item in options:
        if item["id"] == nid or item["label"].casefold() == raw.casefold():
            return item["id"]
    return nid


def _find_lessons_sheet(wb):
    candidates = []
    if SHEET_LESSONS in wb.sheetnames:
        candidates.append(wb[SHEET_LESSONS])
    candidates.extend(ws for ws in wb.worksheets if ws.title not in {SHEET_INSTRUCTIONS, SHEET_SETTINGS, SHEET_META, SHEET_LESSONS})
    last_missing = list(REQUIRED_HEADERS)
    extra = []
    for ws in candidates:
        header_map, extra_headers, missing = _map_headers(ws)
        extra = extra_headers
        if missing:
            last_missing = missing
            continue
        return ws, header_map, extra_headers
    missing = last_missing[0] if last_missing else "Название урока"
    raise PlanExcelError(f"Файл не соответствует шаблону. Не найден столбец «{missing}».")


def _map_headers(ws) -> tuple[dict, list, list]:
    header_map = {}
    extra = []
    for col_idx in range(1, (ws.max_column or 0) + 1):
        title = str(ws.cell(1, col_idx).value or "").strip()
        if not title:
            continue
        key = HEADER_ALIASES.get(title)
        if key:
            header_map[key] = col_idx
        else:
            extra.append(title)
    missing = []
    if "title" not in header_map:
        missing.append("Название урока")
    if "topic" not in header_map:
        missing.append("Тема")
    if "scheduled_date" not in header_map:
        missing.append("Дата")
    return header_map, extra, missing


def _row_is_empty(raw: dict) -> bool:
    for key in ("title", "topic", "subtopic", "task_number", "goal", "description", "homework_description", "teacher_comment"):
        if str(raw.get(key) or "").strip():
            return False
    if raw.get("scheduled_date") not in (None, ""):
        return False
    return True


def _parse_lesson_row(raw: dict, *, excel_row: int, topic_canon: dict) -> dict:
    messages = []
    status_name = "ready"

    title, title_cut = _clip(_cell_text(raw.get("title"), collapse=True), MAX_TITLE)
    topic, topic_cut = _clip(_cell_text(raw.get("topic"), collapse=True), MAX_TOPIC)
    subtopic, sub_cut = _clip(_cell_text(raw.get("subtopic"), collapse=True), MAX_SUBTOPIC)
    task_number, task_cut = _clip(_cell_text(raw.get("task_number"), collapse=True), MAX_TASK_NUMBER)
    goal = _cell_text(raw.get("goal"))
    description = _cell_text(raw.get("description"))
    homework = _cell_text(raw.get("homework_description"))
    comment = _cell_text(raw.get("teacher_comment"))

    if title_cut:
        messages.append("Название обрезано до 255 символов.")
        status_name = "warning"
    if topic_cut:
        messages.append("Тема обрезана до 500 символов.")
        status_name = "warning"
    if sub_cut:
        messages.append("Подтема обрезана до 255 символов.")
        status_name = "warning"
    if task_cut:
        messages.append("№ задания обрезан до 32 символов.")
        status_name = "warning"

    if not title and not topic:
        return {
            "excel_row": excel_row,
            "status": "error",
            "result": "Ошибка",
            "messages": ["Укажите название урока или тему."],
            "item": _empty_item(excel_row, subtopic, task_number, goal, description, homework, comment),
        }

    if not title:
        title = topic

    topic_key = normalize_topic_key(topic)
    if topic_key and topic_key in topic_canon:
        topic = topic_canon[topic_key]
    elif topic:
        topic_canon[topic_key] = topic

    scheduled, date_error = parse_excel_date(raw.get("scheduled_date"))
    if date_error:
        messages.append(f"Строка {excel_row}: {date_error}")
        status_name = "error"
        date_source = ""
    elif scheduled:
        date_source = "manual"
    else:
        date_source = "automatic"

    return {
        "excel_row": excel_row,
        "status": status_name,
        "result": "",
        "messages": messages,
        "item": {
            "id": None,
            "source_item_id": _parse_item_id(raw.get("item_id")),
            "title": title,
            "topic": topic,
            "subtopic": subtopic,
            "task_number": task_number,
            "goal": goal,
            "description": description,
            "homework_description": homework,
            "teacher_comment": comment,
            "scheduled_date": scheduled.isoformat() if scheduled else None,
            "date_source": date_source,
            "order": excel_row - 1,
        },
    }


def _empty_item(excel_row, subtopic, task_number, goal, description, homework, comment):
    return {
        "id": None,
        "source_item_id": None,
        "title": "",
        "topic": "",
        "subtopic": subtopic,
        "task_number": task_number,
        "goal": goal,
        "description": description,
        "homework_description": homework,
        "teacher_comment": comment,
        "scheduled_date": None,
        "date_source": "",
        "order": excel_row - 1,
    }


def _with_canonical_topic(row: dict, topic_canon: dict) -> dict:
    item = dict(row.get("item") or {})
    topic = item.get("topic") or ""
    key = normalize_topic_key(topic)
    if key and key in topic_canon:
        item["topic"] = topic_canon[key]
    elif topic:
        topic_canon[key] = topic
    return {**row, "item": item}


def _cell_text(value: Any, *, collapse: bool = False) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    elif isinstance(value, datetime):
        text = value.date().isoformat()
    elif isinstance(value, date):
        text = value.isoformat()
    else:
        text = str(value).strip()
        if text.startswith("'") and len(text) > 1 and text[1:2] in {"=", "+", "-", "@"}:
            text = text[1:]
    if collapse:
        return " ".join(text.split()) if text else ""
    return text


def _parse_item_id(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        number = int(float(str(value).strip()))
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def _first_explicit_date(rows: list[dict]) -> date | None:
    for row in rows:
        parsed = parse_plan_date((row.get("item") or {}).get("scheduled_date"))
        if parsed:
            return parsed
    return None


def _existing_entries(existing_items: list, interval: str) -> list[dict]:
    rows = []
    for item in existing_items:
        if isinstance(item, dict):
            rows.append({
                "id": item.get("id") or item.get("pk"),
                "title": item.get("title") or "",
                "topic": item.get("topic") or "",
                "subtopic": item.get("subtopic") or "",
                "task_number": item.get("task_number") or "",
                "goal": item.get("goal") or "",
                "description": item.get("description") or "",
                "homework_description": item.get("homework_description") or "",
                "teacher_comment": item.get("teacher_comment") or "",
                "scheduled_date": iso_plan_date(item.get("scheduled_date")) or None,
                "date_source": item.get("date_source") or "",
                "order": item.get("order") or 0,
                "in_use": bool(item.get("in_use")),
            })
        else:
            rows.append({
                "id": item.pk,
                "title": item.title,
                "topic": item.topic,
                "subtopic": item.subtopic,
                "task_number": item.task_number,
                "goal": item.goal,
                "description": item.description,
                "homework_description": item.homework_description,
                "teacher_comment": item.teacher_comment,
                "scheduled_date": iso_plan_date(item.scheduled_date) or None,
                "date_source": "",
                "order": item.order,
                "in_use": _item_in_use(item),
            })
    rows.sort(key=lambda item: (item.get("order") or 0, item.get("id") or 0))
    first = parse_plan_date(rows[0]["scheduled_date"]) if rows else None
    planned = generate_plan_dates(first, len(rows), interval) if first else []
    for index, row in enumerate(rows):
        if row.get("date_source") in {"manual", "automatic"}:
            continue
        current = parse_plan_date(row.get("scheduled_date"))
        auto = planned[index] if index < len(planned) else None
        if current and auto and current != auto and index > 0:
            row["date_source"] = "manual"
        elif current:
            row["date_source"] = "automatic"
        else:
            row["date_source"] = ""
    return rows


def _copy_item(item: dict) -> dict:
    copied = dict(item)
    copied["id"] = None
    return copied


def _resolve_replace_all(rows: list[dict], existing: list[dict], start, interval: str) -> dict:
    prepared = []
    for row in rows:
        item = dict(row["item"])
        item["id"] = None
        prepared.append({**row, "item": item, "result": "Будет в новом плане" if row["status"] != "error" else "Ошибка"})
    _assign_automatic_dates(prepared, start, interval)
    blocking = []
    in_use = [entry for entry in existing if entry.get("in_use")]
    if in_use:
        titles = ", ".join(entry.get("title") or "Без названия" for entry in in_use[:3])
        blocking.append(f"Нельзя обновить план целиком: занятие «{titles}» уже стоит в расписании.")
    operations = []
    order = 1
    for row in prepared:
        if row["status"] == "error":
            continue
        operations.append({"op": "create", "order": order, "item": row["item"]})
        order += 1
    current = len(existing)
    after = sum(1 for row in prepared if row["status"] != "error")
    confirmation = (
        f"В плане сейчас {current} {_lessons_word(current)}. "
        f"После обновления останется {after} {_lessons_word(after)} из Excel. "
        f"Это действие заменит текущий состав плана."
    )
    return {
        "rows": prepared,
        "operations": operations,
        "after_count": after,
        "added": after,
        "replaced": 0,
        "confirmation": confirmation,
        "blocking_errors": blocking,
        "warnings": [],
        "shifts": [],
    }


def _resolve_insert(rows: list[dict], existing: list[dict], start, interval: str) -> dict:
    prepared = []
    inserts = []
    for row in rows:
        item = dict(row["item"])
        item["id"] = None
        status_name = row["status"]
        messages = list(row.get("messages") or [])
        scheduled = parse_plan_date(item.get("scheduled_date"))
        if status_name != "error" and not scheduled:
            status_name = "error"
            messages.append(f"Строка {row['excel_row']}: укажите дату, на которую нужно добавить урок.")
        result = ""
        if status_name == "error":
            result = "Ошибка"
        elif scheduled:
            result = f"{format_ru_date(scheduled)} → будет добавлен новый урок на эту дату"
            item["date_source"] = "manual"
            item["scheduled_date"] = scheduled.isoformat()
        prepared.append({**row, "item": item, "status": status_name, "messages": messages, "result": result})
        if status_name != "error" and scheduled:
            inserts.append(prepared[-1])

    sequence = [dict(entry) for entry in existing]
    inserts_sorted = sorted(
        inserts,
        key=lambda row: (parse_plan_date(row["item"]["scheduled_date"]), row["excel_row"]),
    )
    first_insert_index = None
    for row in inserts_sorted:
        insert_date = parse_plan_date(row["item"]["scheduled_date"])
        index = next(
            (
                pos
                for pos, entry in enumerate(sequence)
                if parse_plan_date(entry.get("scheduled_date")) is None
                or parse_plan_date(entry.get("scheduled_date")) >= insert_date
            ),
            len(sequence),
        )
        new_entry = {
            **_copy_item(row["item"]),
            "kind": "new",
            "excel_row": row["excel_row"],
            "date_source": "manual",
            "scheduled_date": insert_date.isoformat(),
        }
        sequence.insert(index, new_entry)
        first_insert_index = index if first_insert_index is None else min(first_insert_index, index)

    before_dates = {entry.get("id"): entry.get("scheduled_date") for entry in existing if entry.get("id")}
    if sequence and first_insert_index is not None:
        start_for_shift = parse_plan_date(sequence[first_insert_index].get("scheduled_date")) or start
        apply_sequence_dates(sequence, start_for_shift, interval, from_index=first_insert_index, preserve_manual=True)

    warnings = []
    if sequence:
        warnings.extend(_same_date_warnings(sequence))
    shifts = []
    for entry in sequence:
        old_id = entry.get("id")
        if not old_id:
            continue
        new_date = entry.get("scheduled_date")
        old_date = before_dates.get(old_id)
        if old_date != new_date:
            shifts.append({
                "id": old_id,
                "title": entry.get("title") or "",
                "from": old_date,
                "to": new_date,
                "date_source": entry.get("date_source") or "",
            })

    operations = []
    for order, entry in enumerate(sequence, start=1):
        payload = {
            "title": entry.get("title") or "",
            "topic": entry.get("topic") or "",
            "subtopic": entry.get("subtopic") or "",
            "task_number": entry.get("task_number") or "",
            "goal": entry.get("goal") or "",
            "description": entry.get("description") or "",
            "homework_description": entry.get("homework_description") or "",
            "teacher_comment": entry.get("teacher_comment") or "",
            "scheduled_date": entry.get("scheduled_date"),
            "date_source": entry.get("date_source") or "",
        }
        if entry.get("kind") == "new" or not entry.get("id"):
            operations.append({"op": "create", "order": order, "item": payload})
        else:
            operations.append({"op": "update", "id": entry.get("id"), "order": order, "item": payload})

    added = sum(1 for row in prepared if row["status"] != "error")
    confirmation = (
        f"Будет добавлено: {added} {_lessons_word(added)}. "
        "Последующие уроки будут сдвинуты согласно расписанию плана."
    )
    return {
        "rows": prepared,
        "operations": operations,
        "after_count": len(existing) + added,
        "added": added,
        "replaced": 0,
        "confirmation": confirmation,
        "blocking_errors": [],
        "warnings": warnings,
        "shifts": shifts,
    }


def _same_date_warnings(sequence: list[dict]) -> list[str]:
    grouped = {}
    for entry in sequence:
        key = iso_plan_date(entry.get("scheduled_date"))
        if not key:
            continue
        grouped.setdefault(key, []).append(entry)
    warnings = []
    for key, entries in grouped.items():
        if len(entries) < 2:
            continue
        warnings.append(
            f"На {format_ru_date(key)} окажется несколько уроков. Ручные даты не были сдвинуты автоматически."
        )
    return warnings


def _resolve_replace_dates(rows: list[dict], existing: list[dict]) -> dict:
    by_date = {}
    for entry in existing:
        key = iso_plan_date(entry.get("scheduled_date"))
        if not key:
            continue
        by_date.setdefault(key, []).append(entry)

    date_rows = {}
    prepared = []
    operations = []
    for row in rows:
        item = dict(row["item"])
        status_name = row["status"]
        messages = list(row.get("messages") or [])
        scheduled = parse_plan_date(item.get("scheduled_date"))
        current = None
        result = "Ошибка"
        if status_name != "error" and not scheduled:
            status_name = "error"
            messages.append(f"Строка {row['excel_row']}: укажите дату урока, который нужно заменить.")
        if scheduled:
            key = scheduled.isoformat()
            date_rows.setdefault(key, []).append(row["excel_row"])
            matches = by_date.get(key) or []
            if status_name != "error" and not matches:
                status_name = "error"
                messages.append(
                    f"Строка {row['excel_row']}: в текущем плане нет урока на {format_ru_date(scheduled)}. Заменить нечего."
                )
            elif status_name != "error" and len(matches) > 1:
                status_name = "error"
                messages.append(
                    f"Строка {row['excel_row']}: на {format_ru_date(scheduled)} в плане несколько уроков, нельзя однозначно выбрать, какой заменить."
                )
            elif status_name != "error":
                current = matches[0]
                item["id"] = current.get("id")
                item["scheduled_date"] = key
                item["date_source"] = current.get("date_source") or item.get("date_source")
                result = (
                    f"{format_ru_date(scheduled)}\n"
                    f"Было: {current.get('title') or current.get('topic') or '—'}\n"
                    f"Станет: {item.get('title') or item.get('topic') or '—'}"
                )
        prepared.append({
            **row,
            "item": item,
            "status": status_name,
            "messages": messages,
            "result": result,
            "current": {
                "id": current.get("id"),
                "title": current.get("title") or "",
                "topic": current.get("topic") or "",
                "scheduled_date": current.get("scheduled_date"),
            } if current else None,
        })

    for key, excel_rows in date_rows.items():
        if len(excel_rows) < 2:
            continue
        message = f"В файле несколько строк для замены урока на {format_ru_date(key)}."
        for row in prepared:
            if parse_plan_date(row["item"].get("scheduled_date")) and iso_plan_date(row["item"].get("scheduled_date")) == key:
                row["status"] = "error"
                row["result"] = "Ошибка"
                if message not in row["messages"]:
                    row["messages"].append(message)
                row["current"] = None

    replaced = 0
    seen_targets = set()
    for row in prepared:
        if row["status"] == "error" or not row.get("current"):
            continue
        target_id = row["current"]["id"]
        if not target_id or target_id in seen_targets:
            continue
        seen_targets.add(target_id)
        current = next(entry for entry in existing if entry.get("id") == target_id)
        operations.append({
            "op": "update",
            "id": target_id,
            "order": current.get("order") or 1,
            "item": {
                **row["item"],
                "scheduled_date": current.get("scheduled_date"),
                "date_source": current.get("date_source") or "",
            },
        })
        replaced += 1

    return {
        "rows": prepared,
        "operations": operations,
        "after_count": len(existing),
        "added": 0,
        "replaced": replaced,
        "confirmation": f"Будет заменено: {replaced} {_lessons_word(replaced)}. Добавлено: 0.",
        "blocking_errors": [],
        "warnings": [],
        "shifts": [],
    }


def _assign_automatic_dates(rows: list[dict], start: date | None, interval: str) -> None:
    if not start:
        return
    dates = generate_plan_dates(start, len(rows), interval)
    for index, row in enumerate(rows):
        item = row["item"]
        if item.get("date_source") == "manual" and item.get("scheduled_date"):
            continue
        if item.get("scheduled_date"):
            continue
        generated = dates[index] if index < len(dates) else None
        if generated:
            item["scheduled_date"] = generated.isoformat()
            item["date_source"] = "automatic"


def items_from_plan(plan: LessonPlan, interval: str = "weekly") -> list[dict]:
    items = list(plan.items.order_by("order", "id"))
    return _existing_entries(items, interval)

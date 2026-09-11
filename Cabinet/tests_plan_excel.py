from datetime import date, datetime
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from rest_framework.exceptions import ValidationError as DRFValidationError
from rest_framework.test import APIClient

from Cabinet.choices import PlanStatus
from Cabinet.models import LessonPlan, LessonPlanItem, Profile
from Cabinet.plan_dates import generate_plan_dates
from Cabinet.plan_excel import (
    COLUMNS,
    REQUIRED_HEADERS,
    SHEET_INSTRUCTIONS,
    SHEET_LESSONS,
    SHEET_SETTINGS,
    TEMPLATE_VERSION,
    PlanExcelError,
    build_plan_workbook,
    excel_safe_text,
    parse_excel_date,
    parse_plan_workbook,
    read_uploaded_bytes,
)
from Cabinet.serializers import LessonPlanItemEditorSerializer


def _col(key):
    return next(index for index, col in enumerate(COLUMNS, start=1) if col["key"] == key)


def _patch_workbook(content, mutator):
    wb = load_workbook(BytesIO(content))
    mutator(wb)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class PlanExcelHelperTests(TestCase):
    def test_parses_calendar_dates_without_timezone_shift(self):
        self.assertEqual(parse_excel_date("11.09.2026")[0], date(2026, 9, 11))
        self.assertEqual(parse_excel_date("2026-09-11")[0], date(2026, 9, 11))
        self.assertEqual(parse_excel_date(date(2026, 9, 11))[0], date(2026, 9, 11))
        self.assertIsNone(parse_excel_date("31.02.2026")[0])
        self.assertIsNotNone(parse_excel_date("31.02.2026")[1])

    def test_formula_injection_is_prefixed_on_export(self):
        self.assertEqual(excel_safe_text("=1+1"), "'=1+1")
        self.assertEqual(excel_safe_text("Системы счисления"), "Системы счисления")

    def test_template_is_editable_without_password(self):
        content = build_plan_workbook(
            {"title": "Информатика", "subject": "inf", "direction": "ege", "grade": "10–11", "start_date": "2026-09-11"},
            [],
        )
        wb = load_workbook(BytesIO(content))
        self.assertIn(SHEET_LESSONS, wb.sheetnames)
        self.assertIn(SHEET_INSTRUCTIONS, wb.sheetnames)
        self.assertIn(SHEET_SETTINGS, wb.sheetnames)
        self.assertEqual(wb["_meta"]["B1"].value, TEMPLATE_VERSION)
        headers = [cell.value for cell in wb[SHEET_LESSONS][1]]
        for name in REQUIRED_HEADERS:
            self.assertIn(name, headers)
        ws = wb[SHEET_LESSONS]
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.tables)
        self.assertTrue(ws.protection.sheet)
        self.assertIn(ws.protection.password, (None, ""))
        self.assertTrue(ws.protection.insertRows)
        self.assertFalse(ws.protection.insertColumns)
        id_col = next(index for index, col in enumerate(COLUMNS, start=1) if col["key"] == "item_id")
        self.assertTrue(ws.column_dimensions[get_column_letter(id_col)].hidden)
        for index, col in enumerate(COLUMNS, start=1):
            self.assertTrue(ws.cell(1, index).protection.locked)
            if col["hidden"]:
                self.assertTrue(ws.column_dimensions[get_column_letter(index)].hidden)
            if not col["locked"]:
                self.assertFalse(ws.cell(2, index).protection.locked)
        settings = wb[SHEET_SETTINGS]
        self.assertTrue(settings.protection.sheet)
        self.assertIn(settings.protection.password, (None, ""))
        self.assertEqual(settings["A1"].value, "НАСТРОЙКИ ПЛАНА")
        self.assertTrue(settings["A4"].protection.locked)
        self.assertFalse(settings["B4"].protection.locked)
        self.assertFalse(settings["B5"].protection.locked)
        self.assertFalse(settings["B8"].protection.locked)
        instruction = str(wb[SHEET_INSTRUCTIONS]["A1"].value)
        self.assertIn("План занятий", instruction)
        joined = "\n".join(str(cell.value or "") for row in wb[SHEET_INSTRUCTIONS].iter_rows(max_col=5) for cell in row)
        self.assertIn("Для чего нужен файл", joined)
        self.assertIn("Обновить план целиком", joined)
        self.assertIn("Добавить уроки", joined)
        self.assertIn("Заменить уроки по датам", joined)
        self.assertIn("Пошаговая загрузка", joined)


class PlanExcelApiTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(username="excel_teacher", password="pass")
        self.teacher.profile.role = Profile.Role.TEACHER
        self.teacher.profile.save(update_fields=["role"])
        self.plan = LessonPlan.objects.create(
            teacher=self.teacher,
            title="Информатика",
            subject="inf",
            direction="ege",
            status=PlanStatus.DRAFT,
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.teacher)

    def _xlsx(self, content, name="plan.xlsx"):
        return SimpleUploadedFile(
            name,
            content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_download_template(self):
        resp = self.client.get("/api/cabinet/lesson-plans/excel-template/?subject=inf&direction=ege")
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resp["Content-Type"],
        )
        wb = load_workbook(BytesIO(resp.content))
        self.assertIn("Уроки", wb.sheetnames)

    def test_import_ten_lessons_roundtrip(self):
        items = [
            {
                "title": f"Урок {n}",
                "topic": "Системы счисления" if n % 2 else "Кодирование информации",
                "subtopic": "Перевод чисел",
                "task_number": "5",
                "goal": "Повторить тему",
                "description": "Разбор задач",
                "homework_description": "№5",
                "teacher_comment": "Кириллица",
                "scheduled_date": date(2026, 9, 11) if n == 1 else None,
            }
            for n in range(1, 11)
        ]
        content = build_plan_workbook(
            {"start_date": "2026-09-11", "interval": "twice_weekly", "subject": "inf"},
            items,
        )
        resp = self.client.post(
            f"/api/cabinet/lesson-plans/{self.plan.pk}/import-excel/",
            {"file": self._xlsx(content), "mode": "replace_all", "start_date": "2026-09-11", "interval": "twice_weekly"},
            format="multipart",
        )
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertEqual(self.plan.items.count(), 10)
        first = self.plan.items.order_by("order").first()
        self.assertEqual(first.scheduled_date, date(2026, 9, 11))
        self.assertEqual(first.topic, "Системы счисления")

    def test_manual_date_is_kept(self):
        items = [
            {"title": "A", "topic": "Тема", "scheduled_date": None},
            {"title": "B", "topic": "Тема", "scheduled_date": date(2026, 9, 18)},
            {"title": "C", "topic": "Тема", "scheduled_date": None},
        ]
        content = build_plan_workbook({"start_date": "2026-09-11", "interval": "weekly"}, items)
        resp = self.client.post(
            f"/api/cabinet/lesson-plans/{self.plan.pk}/import-excel/",
            {"file": self._xlsx(content), "mode": "replace_all", "start_date": "2026-09-11", "interval": "weekly"},
            format="multipart",
        )
        self.assertEqual(resp.status_code, 200, resp.content)
        dates = list(self.plan.items.order_by("order").values_list("scheduled_date", flat=True))
        self.assertEqual(dates[0], date(2026, 9, 11))
        self.assertEqual(dates[1], date(2026, 9, 18))
        self.assertEqual(dates[2], date(2026, 9, 25))

    def test_missing_header_is_user_error(self):
        content = build_plan_workbook({}, [{"title": "A", "topic": "Тема"}])
        wb = load_workbook(BytesIO(content))
        ws = wb[SHEET_LESSONS]
        for cell in ws[1]:
            if cell.value == "Тема":
                cell.value = "Темы"
                break
        buf = BytesIO()
        wb.save(buf)
        resp = self.client.post(
            "/api/cabinet/lesson-plans/excel-preview/",
            {"file": self._xlsx(buf.getvalue())},
            format="multipart",
        )
        self.assertEqual(resp.status_code, 400)
        self.assertIn("Тема", resp.json()["detail"])

    def test_legacy_date_header_still_works(self):
        content = build_plan_workbook({}, [{"title": "A", "topic": "Тема", "scheduled_date": date(2026, 9, 11)}])
        content = _patch_workbook(content, lambda wb: setattr(wb[SHEET_LESSONS].cell(1, _col("scheduled_date")), "value", "Дата занятия"))
        preview = parse_plan_workbook(content, mode="replace_all")
        self.assertTrue(preview["can_import"])
        self.assertEqual(preview["rows"][0]["item"]["scheduled_date"], "2026-09-11")

    def test_extra_column_is_ignored(self):
        content = build_plan_workbook({}, [{"title": "A", "topic": "Алгебра логики"}])
        wb = load_workbook(BytesIO(content))
        ws = wb[SHEET_LESSONS]
        ws.cell(1, 20, "Мои заметки")
        ws.cell(2, 20, "secret")
        buf = BytesIO()
        wb.save(buf)
        preview = parse_plan_workbook(buf.getvalue())
        self.assertTrue(preview["can_import"])
        self.assertTrue(any("Лишние столбцы" in item for item in preview["warnings"]))

    def test_empty_rows_are_skipped(self):
        items = [
            {"title": "A", "topic": "Тема"},
            {"title": "", "topic": ""},
            {"title": "B", "topic": "Тема"},
        ]
        content = build_plan_workbook({}, items)
        preview = parse_plan_workbook(content)
        self.assertEqual(preview["summary"]["found"], 2)
        self.assertEqual(preview["summary"]["skip"], 1)
        self.assertTrue(all(row["status"] != "skip" for row in preview["rows"]))

    def test_trailing_blank_template_rows_are_not_skipped_warnings(self):
        content = build_plan_workbook({}, [{"title": "A", "topic": "Системы счисления"}])
        preview = parse_plan_workbook(content)
        self.assertEqual(preview["summary"]["found"], 1)
        self.assertEqual(preview["summary"]["skip"], 0)
        self.assertEqual(preview["rows"][0]["status"], "ready")

    def test_task_number_keeps_range_text(self):
        content = build_plan_workbook({}, [{
            "title": "A",
            "topic": "Тема",
            "task_number": "1-5",
        }])
        wb = load_workbook(BytesIO(content))
        ws = wb[SHEET_LESSONS]
        task_col = next(index for index, col in enumerate(COLUMNS, start=1) if col["key"] == "task_number")
        self.assertEqual(ws.cell(2, task_col).number_format, "@")
        self.assertEqual(ws.cell(2, task_col).value, "1-5")
        preview = parse_plan_workbook(content)
        self.assertEqual(preview["rows"][0]["item"]["task_number"], "1-5")

    def test_title_or_topic_is_enough(self):
        content = build_plan_workbook({}, [{"title": "", "topic": "Кодирование текстовой информации"}])
        preview = parse_plan_workbook(content)
        self.assertTrue(preview["can_import"])
        self.assertEqual(preview["rows"][0]["item"]["title"], "Кодирование текстовой информации")


class PlanExcelAcceptanceTests(TestCase):
    def setUp(self):
        self.teacher = User.objects.create_user(username="excel_accept", password="pass")
        self.teacher.profile.role = Profile.Role.TEACHER
        self.teacher.profile.save(update_fields=["role"])
        self.other = User.objects.create_user(username="excel_stranger", password="pass")
        self.other.profile.role = Profile.Role.TEACHER
        self.other.profile.save(update_fields=["role"])
        self.plan = LessonPlan.objects.create(
            teacher=self.teacher,
            title="Информатика",
            subject="inf",
            direction="ege",
            status=PlanStatus.DRAFT,
        )
        self.other_plan = LessonPlan.objects.create(
            teacher=self.other,
            title="Чужой план",
            subject="inf",
            direction="ege",
            status=PlanStatus.DRAFT,
        )
        self.own_other_plan = LessonPlan.objects.create(
            teacher=self.teacher,
            title="Другой свой план",
            subject="inf",
            direction="ege",
            status=PlanStatus.DRAFT,
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.teacher)

    def _xlsx(self, content, name="plan.xlsx"):
        return SimpleUploadedFile(
            name,
            content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def _import(self, content, *, mode="replace_all", start_date="2026-09-11", interval="weekly", plan=None):
        plan = plan or self.plan
        payload = {"file": self._xlsx(content), "mode": mode}
        if start_date:
            payload["start_date"] = start_date
        if interval:
            payload["interval"] = interval
        return self.client.post(
            f"/api/cabinet/lesson-plans/{plan.pk}/import-excel/",
            payload,
            format="multipart",
        )

    def _preview(self, content, *, plan=None, mode="replace_all", start_date="2026-09-11", interval="weekly"):
        payload = {"file": self._xlsx(content), "mode": mode}
        if plan is not None:
            payload["plan_id"] = plan.pk
        if start_date:
            payload["start_date"] = start_date
        if interval:
            payload["interval"] = interval
        return self.client.post("/api/cabinet/lesson-plans/excel-preview/", payload, format="multipart")

    def _reload(self, plan=None):
        plan = plan or self.plan
        return list(plan.items.order_by("order", "id").values("id", "title", "topic", "scheduled_date", "order"))

    def test_excel_date_and_text_date_stay_on_calendar_day(self):
        self.assertEqual(parse_excel_date("11.09.2026")[0], date(2026, 9, 11))
        self.assertEqual(parse_excel_date(date(2026, 9, 11))[0], date(2026, 9, 11))
        self.assertEqual(parse_excel_date(datetime(2026, 9, 11, 0, 0, 0))[0], date(2026, 9, 11))
        self.assertEqual(parse_excel_date(to_excel(date(2026, 9, 11)))[0], date(2026, 9, 11))

        content = build_plan_workbook({}, [{"title": "Дата", "topic": "Тема"}])
        content = _patch_workbook(content, lambda wb: wb[SHEET_LESSONS].cell(
            2, _col("scheduled_date"), datetime(2026, 9, 11, 0, 0, 0)
        ))
        resp = self._import(content, mode="replace_all")
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertEqual(self.plan.items.get().scheduled_date, date(2026, 9, 11))

    def test_replace_all_twenty_with_eight_keeps_excel_order(self):
        for n in range(1, 21):
            LessonPlanItem.objects.create(plan=self.plan, order=n, title=f"Старый {n}", topic="Тема")
        titles = [f"Новый {n}" for n in range(1, 9)]
        content = build_plan_workbook({}, [{"title": title, "topic": "Тема"} for title in titles])
        preview = self._preview(content, plan=self.plan, mode="replace_all").json()
        self.assertTrue(preview["can_import"])
        self.assertEqual(preview["summary"]["current_count"], 20)
        self.assertEqual(preview["summary"]["after_count"], 8)
        resp = self._import(content, mode="replace_all")
        self.assertEqual(resp.status_code, 200, resp.content)
        actual = list(self.plan.items.order_by("order").values_list("title", flat=True))
        self.assertEqual(actual, titles)
        self.assertEqual(self._reload(), self._reload())

    def test_replace_all_rolls_back_on_row_error(self):
        for n in range(1, 21):
            LessonPlanItem.objects.create(plan=self.plan, order=n, title=f"Старый {n}", topic="Тема")
        items = [{"title": f"Новый {n}", "topic": "Тема"} for n in range(1, 9)]
        content = build_plan_workbook({}, items)

        def break_fifth_row(wb):
            ws = wb[SHEET_LESSONS]
            ws.cell(6, _col("title")).value = ""
            ws.cell(6, _col("topic")).value = ""
            ws.cell(6, _col("scheduled_date")).value = date(2026, 9, 11)

        content = _patch_workbook(content, break_fifth_row)
        preview = self._preview(content, plan=self.plan, mode="replace_all").json()
        self.assertFalse(preview["can_import"])
        resp = self._import(content, mode="replace_all")
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(self.plan.items.count(), 20)
        self.assertEqual(
            list(self.plan.items.order_by("order").values_list("title", flat=True))[:3],
            ["Старый 1", "Старый 2", "Старый 3"],
        )

    def test_replace_all_does_not_match_by_title(self):
        item = LessonPlanItem.objects.create(plan=self.plan, order=1, title="Урок 1", topic="Тема")
        content = build_plan_workbook({}, [{"title": "Урок 1", "topic": "Другая"}])
        resp = self._import(content, mode="replace_all")
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertEqual(self.plan.items.count(), 1)
        self.assertFalse(LessonPlanItem.objects.filter(pk=item.pk).exists())
        self.assertEqual(self.plan.items.get().topic, "Другая")

    def test_insert_shifts_subsequent_lessons_with_existing_scheduler(self):
        LessonPlanItem.objects.create(plan=self.plan, order=1, title="Урок A", topic="Тема", scheduled_date=date(2026, 9, 11))
        LessonPlanItem.objects.create(plan=self.plan, order=2, title="Урок B", topic="Тема", scheduled_date=date(2026, 9, 18))
        LessonPlanItem.objects.create(plan=self.plan, order=3, title="Урок C", topic="Тема", scheduled_date=date(2026, 9, 25))
        content = build_plan_workbook({}, [{"title": "Новый урок", "topic": "Тема", "scheduled_date": date(2026, 9, 18)}])
        with patch("Cabinet.plan_dates.generate_plan_dates", wraps=generate_plan_dates) as spy:
            preview = self._preview(content, plan=self.plan, mode="insert", interval="weekly").json()
            resp = self._import(content, mode="insert", interval="weekly")
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertTrue(spy.called)
        self.assertTrue(preview["can_import"])
        self.assertEqual(preview["summary"]["added"], 1)
        rows = list(self.plan.items.order_by("order"))
        self.assertEqual([item.title for item in rows], ["Урок A", "Новый урок", "Урок B", "Урок C"])
        self.assertEqual([item.scheduled_date for item in rows], [
            date(2026, 9, 11),
            date(2026, 9, 18),
            date(2026, 9, 25),
            date(2026, 10, 2),
        ])
        self.assertEqual(self._reload(), self._reload())

    def test_insert_keeps_manual_dates(self):
        LessonPlanItem.objects.create(plan=self.plan, order=1, title="Урок A", topic="Тема", scheduled_date=date(2026, 9, 11))
        LessonPlanItem.objects.create(plan=self.plan, order=2, title="Урок B", topic="Тема", scheduled_date=date(2026, 9, 18))
        LessonPlanItem.objects.create(plan=self.plan, order=3, title="Урок C", topic="Тема", scheduled_date=date(2026, 10, 2))
        content = build_plan_workbook({}, [{"title": "Новый урок", "topic": "Тема", "scheduled_date": date(2026, 9, 18)}])
        resp = self._import(content, mode="insert", interval="weekly")
        self.assertEqual(resp.status_code, 200, resp.content)
        rows = list(self.plan.items.order_by("order"))
        self.assertEqual([item.title for item in rows], ["Урок A", "Новый урок", "Урок B", "Урок C"])
        self.assertEqual(rows[3].scheduled_date, date(2026, 10, 2))
        sources = {row["id"]: row["date_source"] for row in resp.json()["item_dates"]}
        self.assertEqual(sources[rows[1].pk], "manual")
        self.assertEqual(sources[rows[3].pk], "manual")

    def test_insert_ignores_ids_and_always_creates(self):
        local = LessonPlanItem.objects.create(plan=self.plan, order=1, title="Старый", topic="Тема", scheduled_date=date(2026, 9, 11))
        foreign = LessonPlanItem.objects.create(plan=self.other_plan, order=1, title="Чужой", topic="Секреты", scheduled_date=date(2026, 9, 11))
        content = build_plan_workbook({}, [
            {"id": foreign.pk, "title": "Добавленный из чужого ID", "topic": "Тема", "scheduled_date": date(2026, 9, 18)},
            {"id": local.pk, "title": "Тоже новый", "topic": "Тема", "scheduled_date": date(2026, 9, 25)},
        ])
        resp = self._import(content, mode="insert")
        self.assertEqual(resp.status_code, 200, resp.content)
        local.refresh_from_db()
        foreign.refresh_from_db()
        self.assertEqual(local.title, "Старый")
        self.assertEqual(foreign.title, "Чужой")
        self.assertEqual(self.plan.items.count(), 3)
        self.assertTrue(self.plan.items.filter(title="Добавленный из чужого ID").exists())
        self.assertTrue(self.plan.items.filter(title="Тоже новый").exists())

    def test_replace_dates_updates_existing_day(self):
        a = LessonPlanItem.objects.create(plan=self.plan, order=1, title="Системы счисления", topic="Тема", scheduled_date=date(2026, 9, 11))
        b = LessonPlanItem.objects.create(plan=self.plan, order=2, title="Логика", topic="Тема", scheduled_date=date(2026, 9, 16))
        c = LessonPlanItem.objects.create(plan=self.plan, order=3, title="Кодирование", topic="Тема", scheduled_date=date(2026, 9, 18))
        content = build_plan_workbook({}, [{"title": "Алгебра логики", "topic": "Тема", "scheduled_date": date(2026, 9, 16)}])
        preview = self._preview(content, plan=self.plan, mode="replace_dates").json()
        self.assertTrue(preview["can_import"])
        self.assertIn("Было: Логика", preview["rows"][0]["result"])
        self.assertIn("Станет: Алгебра логики", preview["rows"][0]["result"])
        resp = self._import(content, mode="replace_dates")
        self.assertEqual(resp.status_code, 200, resp.content)
        self.assertEqual(self.plan.items.count(), 3)
        a.refresh_from_db()
        b.refresh_from_db()
        c.refresh_from_db()
        self.assertEqual(a.title, "Системы счисления")
        self.assertEqual(b.title, "Алгебра логики")
        self.assertEqual(c.title, "Кодирование")
        self.assertEqual(b.scheduled_date, date(2026, 9, 16))
        self.assertEqual(b.order, 2)

    def test_replace_dates_missing_day_is_error(self):
        LessonPlanItem.objects.create(plan=self.plan, order=1, title="Есть", topic="Тема", scheduled_date=date(2026, 9, 11))
        content = build_plan_workbook({}, [{"title": "Нет цели", "topic": "Тема", "scheduled_date": date(2026, 9, 23)}])
        preview = self._preview(content, plan=self.plan, mode="replace_dates").json()
        self.assertFalse(preview["can_import"])
        self.assertIn("нет урока на 23.09.2026", preview["rows"][0]["messages"][0])
        resp = self._import(content, mode="replace_dates")
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(self.plan.items.count(), 1)
        self.assertEqual(self.plan.items.get().title, "Есть")

    def test_replace_dates_duplicate_excel_rows_are_error(self):
        LessonPlanItem.objects.create(plan=self.plan, order=1, title="Логика", topic="Тема", scheduled_date=date(2026, 9, 16))
        content = build_plan_workbook({}, [
            {"title": "Первая", "topic": "Тема", "scheduled_date": date(2026, 9, 16)},
            {"title": "Вторая", "topic": "Тема", "scheduled_date": date(2026, 9, 16)},
        ])
        preview = self._preview(content, plan=self.plan, mode="replace_dates").json()
        self.assertFalse(preview["can_import"])
        self.assertTrue(any("несколько строк для замены урока на 16.09.2026" in msg for row in preview["rows"] for msg in row["messages"]))
        resp = self._import(content, mode="replace_dates")
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(self.plan.items.get().title, "Логика")

    def test_replace_dates_ambiguous_existing_lessons_are_error(self):
        LessonPlanItem.objects.create(plan=self.plan, order=1, title="Первый", topic="Тема", scheduled_date=date(2026, 9, 16))
        LessonPlanItem.objects.create(plan=self.plan, order=2, title="Второй", topic="Тема", scheduled_date=date(2026, 9, 16))
        content = build_plan_workbook({}, [{"title": "Замена", "topic": "Тема", "scheduled_date": date(2026, 9, 16)}])
        preview = self._preview(content, plan=self.plan, mode="replace_dates").json()
        self.assertFalse(preview["can_import"])
        self.assertIn("несколько уроков", preview["rows"][0]["messages"][0])
        resp = self._import(content, mode="replace_dates")
        self.assertEqual(resp.status_code, 400)
        self.assertEqual(self.plan.items.count(), 2)

    def test_foreign_id_never_updates_another_plan(self):
        foreign = LessonPlanItem.objects.create(plan=self.other_plan, order=1, title="Чужой", topic="Секреты", scheduled_date=date(2026, 9, 11))
        own_other = LessonPlanItem.objects.create(plan=self.own_other_plan, order=1, title="Другой план", topic="Тема", scheduled_date=date(2026, 9, 11))
        content = build_plan_workbook({}, [
            {"id": foreign.pk, "title": "Попытка захвата", "topic": "Тема", "scheduled_date": date(2026, 9, 11)},
        ])
        for mode in ("replace_all", "insert", "replace_dates"):
            LessonPlanItem.objects.filter(plan=self.plan).delete()
            LessonPlanItem.objects.create(plan=self.plan, order=1, title="Свой", topic="Тема", scheduled_date=date(2026, 9, 11))
            resp = self._import(content, mode=mode)
            self.assertIn(resp.status_code, {200, 400}, resp.content)
            foreign.refresh_from_db()
            own_other.refresh_from_db()
            self.assertEqual(foreign.title, "Чужой")
            self.assertEqual(own_other.title, "Другой план")
            self.assertFalse(self.other_plan.items.filter(title="Попытка захвата").exists())
            self.assertFalse(self.own_other_plan.items.filter(title="Попытка захвата").exists())

    def test_preview_matches_applied_replace_all(self):
        for n in range(1, 4):
            LessonPlanItem.objects.create(plan=self.plan, order=n, title=f"Старый {n}", topic="Тема")
        content = build_plan_workbook({}, [
            {"title": "A", "topic": "Тема"},
            {"title": "B", "topic": "Тема"},
        ])
        preview = self._preview(content, plan=self.plan, mode="replace_all").json()
        applied = self._import(content, mode="replace_all")
        self.assertEqual(applied.status_code, 200, applied.content)
        titles = list(self.plan.items.order_by("order").values_list("title", flat=True))
        self.assertEqual(preview["summary"]["after_count"], len(titles))
        self.assertEqual(titles, ["A", "B"])
        self.assertEqual(preview["rows"][0]["result"], "Будет в новом плане")

    def test_transaction_rolls_back_when_a_later_row_fails(self):
        for n in range(1, 4):
            LessonPlanItem.objects.create(plan=self.plan, order=n, title=f"Старый {n}", topic="Тема")
        content = build_plan_workbook({}, [
            {"title": f"Новый {n}", "topic": "Тема"} for n in range(1, 16)
        ])
        real_save = LessonPlanItemEditorSerializer.save
        calls = {"n": 0}

        def wrapped(self, **kwargs):
            calls["n"] += 1
            if calls["n"] == 15:
                raise DRFValidationError({"title": ["fail"]})
            return real_save(self, **kwargs)

        with patch.object(LessonPlanItemEditorSerializer, "save", wrapped):
            resp = self._import(content, mode="replace_all")
        self.assertEqual(resp.status_code, 400, resp.content)
        self.assertEqual(self.plan.items.count(), 3)
        self.assertEqual(
            list(self.plan.items.order_by("order").values_list("title", flat=True)),
            ["Старый 1", "Старый 2", "Старый 3"],
        )

    def test_mixed_auto_and_manual_dates(self):
        items = [
            {"title": "1", "topic": "Тема", "scheduled_date": None},
            {"title": "2", "topic": "Тема", "scheduled_date": None},
            {"title": "3", "topic": "Тема", "scheduled_date": date(2026, 10, 2)},
            {"title": "4", "topic": "Тема", "scheduled_date": None},
            {"title": "5", "topic": "Тема", "scheduled_date": None},
        ]
        content = build_plan_workbook({"start_date": "2026-09-11", "interval": "weekly"}, items)
        resp = self._import(content, mode="replace_all", start_date="2026-09-11", interval="weekly")
        self.assertEqual(resp.status_code, 200, resp.content)
        dates = list(self.plan.items.order_by("order").values_list("scheduled_date", flat=True))
        self.assertEqual(dates, [
            date(2026, 9, 11),
            date(2026, 9, 18),
            date(2026, 10, 2),
            date(2026, 10, 2),
            date(2026, 10, 9),
        ])
        sources = {row["id"]: row["date_source"] for row in resp.json()["item_dates"]}
        items = list(self.plan.items.order_by("order"))
        self.assertEqual(sources[items[0].pk], "automatic")
        self.assertEqual(sources[items[1].pk], "automatic")
        self.assertEqual(sources[items[2].pk], "manual")
        self.assertEqual(sources[items[3].pk], "automatic")
        self.assertEqual(sources[items[4].pk], "automatic")

    def test_topic_normalizes_without_fuzzy_match(self):
        LessonPlanItem.objects.create(plan=self.plan, order=1, title="База", topic="Системы счисления")
        items = [
            {"title": "A", "topic": "системы   счисления"},
            {"title": "B", "topic": "СИСТЕМЫ СЧИСЛЕНИЯ"},
            {"title": "C", "topic": "Системы счисл"},
        ]
        content = build_plan_workbook({}, items)
        preview = parse_plan_workbook(content, existing_items=list(self.plan.items.all()), mode="insert")
        topics = [row["item"]["topic"] for row in preview["rows"]]
        self.assertEqual(topics[0], "Системы счисления")
        self.assertEqual(topics[1], "Системы счисления")
        self.assertEqual(topics[2], "Системы счисл")

    def test_broken_files_return_user_errors(self):
        missing = _patch_workbook(
            build_plan_workbook({}, [{"title": "A", "topic": "Тема"}]),
            lambda wb: setattr(wb[SHEET_LESSONS].cell(1, _col("title")), "value", "Имя урока"),
        )
        resp = self._preview(missing)
        self.assertEqual(resp.status_code, 400)
        self.assertNotIn("Traceback", resp.content.decode())
        self.assertIn("Название урока", resp.json()["detail"])

        empty = self.client.post(
            "/api/cabinet/lesson-plans/excel-preview/",
            {"file": self._xlsx(b"", name="empty.xlsx")},
            format="multipart",
        )
        self.assertEqual(empty.status_code, 400)
        self.assertIn("пустой", empty.json()["detail"].lower())

        xlsm = self.client.post(
            "/api/cabinet/lesson-plans/excel-preview/",
            {"file": self._xlsx(build_plan_workbook({}, [{"title": "A", "topic": "Тема"}]), name="plan.xlsm")},
            format="multipart",
        )
        self.assertEqual(xlsm.status_code, 400)
        self.assertIn("xlsx", xlsm.json()["detail"].lower())

        xls = self.client.post(
            "/api/cabinet/lesson-plans/excel-preview/",
            {"file": self._xlsx(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", name="plan.xlsx")},
            format="multipart",
        )
        self.assertEqual(xls.status_code, 400)

        other = Workbook()
        other.active["A1"] = "hello"
        buf = BytesIO()
        other.save(buf)
        other_resp = self._preview(buf.getvalue())
        self.assertEqual(other_resp.status_code, 400)

        corrupt = self.client.post(
            "/api/cabinet/lesson-plans/excel-preview/",
            {"file": self._xlsx(b"PK\x03\x04not-an-xlsx")},
            format="multipart",
        )
        self.assertEqual(corrupt.status_code, 400)
        self.assertNotIn("Traceback", corrupt.content.decode())

        class Dummy:
            name = "plan.xlsm"
            size = 10
            def read(self):
                return b"not used"
        with self.assertRaises(PlanExcelError):
            read_uploaded_bytes(Dummy())

    def test_formula_looking_text_roundtrip(self):
        content = build_plan_workbook({}, [{
            "title": "=SUM(A1:A2)",
            "topic": "+123",
            "subtopic": "@test",
            "goal": "-something",
        }])
        wb = load_workbook(BytesIO(content))
        title_cell = wb[SHEET_LESSONS].cell(2, _col("title"))
        self.assertNotEqual(title_cell.data_type, "f")
        self.assertEqual(title_cell.value, "'=SUM(A1:A2)")
        preview = parse_plan_workbook(content)
        item = preview["rows"][0]["item"]
        self.assertEqual(item["title"], "=SUM(A1:A2)")
        self.assertEqual(item["topic"], "+123")
        self.assertEqual(item["subtopic"], "@test")
        self.assertEqual(item["goal"], "-something")

    def test_large_file_keeps_order(self):
        items = [{"title": f"Урок {n:03d}", "topic": "Тема"} for n in range(1, 151)]
        content = build_plan_workbook({}, items)
        preview = parse_plan_workbook(content)
        self.assertEqual(preview["summary"]["found"], 150)
        self.assertTrue(preview["can_import"])
        resp = self._import(content, mode="replace_all")
        self.assertEqual(resp.status_code, 200, resp.content)
        titles = list(self.plan.items.order_by("order").values_list("title", flat=True))
        self.assertEqual(titles[0], "Урок 001")
        self.assertEqual(titles[-1], "Урок 150")
        self.assertEqual(len(titles), 150)
